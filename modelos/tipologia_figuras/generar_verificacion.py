"""Hoja de Excel para comprobar a mano la tipología de figuras: 6 figuras por tipología, al azar (semilla fija).
Preguntas: ¿la tipología es correcta?, nº real de paneles y, si es espectro, ¿hay curvas superpuestas?
Uso: python modelos/tipologia_figuras/generar_verificacion.py   → Verificar_tipologia.xlsx
"""
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

RAIZ = Path(__file__).resolve().parents[2]
OUT = Path(__file__).resolve().parent
t = pd.read_csv(OUT / "tipologia_figuras.csv", dtype=str, encoding="utf-8-sig")
t = t[t.tipologia != "sin_micro_ni_espectro"]


def ruta(r):
    for base in ("producto/imagenes", "producto/imagenes_v02"):
        p = RAIZ / base / r.pmcid / r.archivo_grafico
        if p.exists():
            return f"../../{base}/{r.pmcid}/{r.archivo_grafico}"
    return ""


m = pd.concat([g.sample(min(6, len(g)), random_state=20261023) for _, g in t.groupby("tipologia")]).sample(frac=1, random_state=7).reset_index(drop=True)
wb = Workbook(); ws = wb.active; ws.title = "Verificar"
cab = ["id", "pmcid", "figura", "tipología (estimada)", "paneles (estimados)", "pie de figura", "imagen",
       "¿La tipología es correcta? (sí/no/dudoso)", "Nº real de paneles", "Tipología correcta (si 'no')", "¿Espectros con curvas superpuestas? (sí/no/no aplica)", "Nota"]
ws.append(cab)
for c in ws[1]:
    c.font = Font(bold=True); c.fill = PatternFill("solid", fgColor="DCEBEA"); c.alignment = Alignment(wrap_text=True, vertical="top")
for i, r in enumerate(m.itertuples(), 1):
    img = ruta(r)
    ws.append([f"TIP-{i:03d}", r.pmcid, r.figura, r.tipologia, r.n_paneles_pie, r.pie, "abrir imagen" if img else "(sin imagen local)", "", "", "", "", ""])
    ws.cell(ws.max_row, 7).hyperlink = img or f"https://pmc.ncbi.nlm.nih.gov/articles/{r.pmcid}/"
    ws.cell(ws.max_row, 7).font = Font(color="0563C1", underline="single")
    if not img:
        ws.cell(ws.max_row, 7).value = "abrir artículo (imagen aún no descargada)"
dv = DataValidation(type="list", formula1='"sí,no,dudoso"', allow_blank=True); ws.add_data_validation(dv); dv.add(f"H2:H{ws.max_row}")
dv2 = DataValidation(type="list", formula1='"sí,no,no aplica"', allow_blank=True); ws.add_data_validation(dv2); dv2.add(f"K2:K{ws.max_row}")
dv3 = DataValidation(type="list", formula1='"micro_unica,espectro_unico,espectro_superpuesto,pareja_micro_espectro,pareja_micro,pareja_espectro,compuesta_micro,compuesta_espectro,compuesta_mixta,compuesta_con_otros,una_util_con_otros,micro_multiple_sin_etiquetas,solo_descriptores,sin_micro_ni_espectro"', allow_blank=True)
ws.add_data_validation(dv3); dv3.add(f"J2:J{ws.max_row}")
for col, w in zip("ABCDEFGHIJKL", [9, 13, 11, 24, 10, 80, 14, 18, 10, 26, 22, 30]):
    ws.column_dimensions[col].width = w
for row in ws.iter_rows(min_row=2):
    for c in row:
        c.alignment = Alignment(wrap_text=True, vertical="top")
ws.freeze_panes = "D2"
wb.save(OUT / "Verificar_tipologia.xlsx")
print(len(m), "figuras;", m.tipologia.value_counts().to_dict())
print("con imagen local:", int((m.apply(ruta, axis=1) != "").sum()))
