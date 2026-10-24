"""Genera la auditoría humana (muestra aleatoria estratificada de la base v0.1) en tres libros de Excel, uno por revisor.

Muestra (semilla fija): 120 filas de descriptores, 90 de espectros y 60 micrografías, estratificadas por clase de portador, MÁS un subconjunto común
(15 descriptores, 10 espectros, 6 micrografías) que revisan los tres para medir el acuerdo entre personas.
No modifica nada de curation/. Salida: Auditoria_Hector.xlsx, Auditoria_Bryan.xlsx, Auditoria_Carlos.xlsx, muestra_maestra.csv.
Uso: python modelos/auditoria_humana/generar_muestra.py
"""
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

RAIZ = Path(__file__).resolve().parents[2]
OUT = Path(__file__).resolve().parent
SEMILLA = 20261023
N = {"descriptores": 120, "espectros": 90, "micrografias": 60}
COMUN = {"descriptores": 15, "espectros": 10, "micrografias": 6}
REVISORES = ["Hector", "Bryan", "Carlos"]
rd = lambda n: pd.read_csv(RAIZ / f"curation/{n}.csv", encoding="utf-8-sig")
art = rd("ext_articulos").set_index("pmcid")
tm = pd.read_csv(RAIZ / "curation/v0.1/tabla_modelo.csv", encoding="utf-8-sig").drop_duplicates("pmcid").set_index("pmcid").clase_modelo
tablas = {"descriptores": rd("ext_descriptores"), "espectros": rd("ext_espectros"), "micrografias": rd("ext_paneles")}
rng = np.random.default_rng(SEMILLA)


def estratificada(df, n, excluidos=()):
    df = df.copy()
    df["clase"] = df.pmcid.map(tm).fillna(df.pmcid.map(art.carrier_type)).fillna("otro")
    df = df[~df.index.isin(excluidos)]
    sizes = df.clase.value_counts()
    asign = (sizes / sizes.sum() * n).round().astype(int).clip(lower=min(8, n // 8))
    while asign.sum() > n:                       # recorta de las clases mayores
        asign[asign.idxmax()] -= 1
    while asign.sum() < n:
        asign[asign.idxmax()] += 1
    partes = [df[df.clase == c].sample(min(k, int((df.clase == c).sum())), random_state=int(rng.integers(1e9))) for c, k in asign.items()]
    return pd.concat(partes)


COLS = {
    "descriptores": ["sample", "size_nm", "pdi", "zeta_mV", "medium", "source", "quote"],
    "espectros": ["technique", "fig", "applies_to", "sample", "quote"],
    "micrografias": ["fig", "variant", "sample", "scale_bar_in_caption", "image_file", "quote"],
}
PREG = {
    "descriptores": [("¿La cita respalda esos valores para esa muestra? (sí/no/dudoso)", "ok_cita"), ("¿Necesitaste abrir el artículo (tabla)? (sí/no)", "abrio_articulo"), ("Nota", "nota")],
    "espectros": [("¿Esa figura muestra esa técnica para esa muestra? (sí/no/dudoso)", "ok_figura"), ("Nota", "nota")],
    "micrografias": [("¿Es una micrografía de nanopartículas del artículo? (sí/no/dudoso)", "ok_micrografia"), ("¿La técnica/variante es correcta? (sí/no/dudoso)", "ok_variante"), ("Nota", "nota")],
}
maestra, comunes_ids = [], {}
elegidas = {}
for k in N:
    df = tablas[k].reset_index(drop=True)
    comun = estratificada(df, COMUN[k])
    unico = estratificada(df, N[k], excluidos=comun.index)
    comun = comun.assign(grupo="COMUN")
    unico = unico.sample(frac=1, random_state=int(rng.integers(1e9))).assign(grupo=[REVISORES[i % 3] for i in range(len(unico))])
    s = pd.concat([comun, unico])
    s["id"] = [f"{k[:3].upper()}-{i+1:03d}" for i in range(len(s))]
    elegidas[k] = s
    maestra.append(s.assign(hoja=k))
pd.concat(maestra).to_csv(OUT / "muestra_maestra.csv", index=False, encoding="utf-8-sig")

azul = PatternFill("solid", fgColor="DCEBEA"); amar = PatternFill("solid", fgColor="FFF5D6")
for rev in REVISORES:
    wb = Workbook()
    ws0 = wb.active; ws0.title = "LEEME"
    texto = [
        f"Auditoría humana de la base v0.1 — revisor: {rev}",
        "",
        "Qué se hace: comprobar a mano una muestra aleatoria de lo que extrajo la IA. No se corrige nada: solo se responde sí / no / dudoso.",
        "Cada hoja tiene las filas que te tocan (grupo = tu nombre) y las filas COMUN, que revisan las tres personas para medir el acuerdo. No comentes tus respuestas con los demás hasta terminar.",
        "",
        "Descriptores: lee la cita (columna 'quote'). ¿Respalda ese tamaño/PDI/zeta para esa muestra? Si la cita es una fila de tabla («Muestra | 120 | 0,2 | −20») y no sabes qué columna es cuál, abre el artículo (enlace DOI), busca la tabla y marca 'sí' en '¿Necesitaste abrir el artículo?'.",
        "Espectros: abre la imagen (enlace) o el artículo. ¿La figura muestra esa técnica (FTIR, XRD…) para esa muestra?",
        "Micrografías: abre la imagen (enlace). ¿Es una micrografía (TEM/SEM…) de nanopartículas del artículo y la variante indicada es correcta?",
        "Los enlaces a imágenes son relativos: funcionan si este libro está en modelos/auditoria_humana/ dentro del repositorio y las imágenes en producto/imagenes/.",
        "dudoso = no puedes decidir con lo que ves. No lo fuerces a sí o no.",
        "Cuando termines, guarda y envía el libro a Héctor.",
    ]
    for i, t in enumerate(texto, 1):
        ws0.cell(i, 1, t).alignment = Alignment(wrap_text=True, vertical="top")
    ws0["A1"].font = Font(bold=True, size=13); ws0.column_dimensions["A"].width = 120
    for k in N:
        s = elegidas[k]; s = s[s.grupo.isin([rev, "COMUN"])]
        ws = wb.create_sheet(k)
        cab = ["id", "grupo", "pmcid", "clase", "DOI"] + COLS[k] + (["imagen"] if k != "descriptores" else []) + [p[0] for p in PREG[k]]
        ws.append(cab)
        for c in ws[1]:
            c.font = Font(bold=True); c.fill = azul; c.alignment = Alignment(wrap_text=True, vertical="top")
        for r in s.itertuples():
            fila = [r.id, r.grupo, r.pmcid, r.clase, "abrir artículo"] + [getattr(r, c) if not (isinstance(getattr(r, c), float) and np.isnan(getattr(r, c))) else "" for c in COLS[k]]
            if k != "descriptores":
                if k == "micrografias":
                    img = f"../../producto/imagenes/{r.pmcid}/{r.image_file}"
                else:
                    img = ""
                fila.append("abrir imagen" if img else "(ver artículo)")
            fila += [""] * len(PREG[k])
            ws.append(fila)
            rr = ws.max_row
            ws.cell(rr, 5).hyperlink = r.doi_link; ws.cell(rr, 5).font = Font(color="0563C1", underline="single")
            if k == "micrografias":
                ic = 5 + len(COLS[k]) + 1
                ws.cell(rr, ic).hyperlink = img; ws.cell(rr, ic).font = Font(color="0563C1", underline="single")
            if r.grupo == "COMUN":
                for c in range(1, 3):
                    ws.cell(rr, c).fill = amar
        ncol = len(cab); resp0 = ncol - len(PREG[k]) + 1
        dv = DataValidation(type="list", formula1='"sí,no,dudoso"', allow_blank=True)
        ws.add_data_validation(dv)
        for j, (_, clave) in enumerate(PREG[k]):
            if clave != "nota":
                letra = ws.cell(1, resp0 + j).column_letter
                dv.add(f"{letra}2:{letra}{ws.max_row}")
        for i, w in enumerate([9, 9, 13, 12, 14], 1):
            ws.column_dimensions[ws.cell(1, i).column_letter].width = w
        for i in range(6, ncol + 1):
            ws.column_dimensions[ws.cell(1, i).column_letter].width = 60 if "quote" in str(ws.cell(1, i).value) else 22
        for row in ws.iter_rows(min_row=2):
            for c in row:
                c.alignment = Alignment(wrap_text=True, vertical="top")
        ws.freeze_panes = "F2"
    wb.save(OUT / f"Auditoria_{rev}.xlsx")
    print(rev, {k: int(elegidas[k].grupo.isin([rev, "COMUN"]).sum()) for k in N})
print("maestra:", {k: len(v) for k, v in elegidas.items()})
